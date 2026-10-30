from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from pathlib import PurePath
from typing import Any

from app.ingest.blocks import Block, BlockKind, ParsedDocument, heading_level, lines_to_blocks
from app.ingest.ocr import OcrEngine, OcrUnavailableError
from app.observability import get_logger

log = get_logger(__name__)

MIN_PAGE_TEXT_CHARS = 25
MAX_OCR_PAGES = 300
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100
MAX_DOCX_ENTRIES = 10_000
MAX_IMAGE_PIXELS = 50_000_000

PDF_MIMES = frozenset({"application/pdf", "application/x-pdf"})
DOCX_MIMES = frozenset({"application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
TEXT_MIMES = frozenset({"text/plain", "text/csv"})
MARKDOWN_MIMES = frozenset({"text/markdown", "text/x-markdown"})
HTML_MIMES = frozenset({"text/html", "application/xhtml+xml"})
IMAGE_MIMES = frozenset({"image/png", "image/jpeg", "image/jpg", "image/tiff", "image/bmp", "image/gif", "image/webp"})
EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".csv": "txt",
    ".md": "md",
    ".markdown": "md",
    ".html": "html",
    ".htm": "html",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".tif": "image",
    ".tiff": "image",
    ".bmp": "image",
    ".gif": "image",
    ".webp": "image",
}
SUPPORTED_MIMES = sorted(PDF_MIMES | DOCX_MIMES | TEXT_MIMES | MARKDOWN_MIMES | HTML_MIMES | IMAGE_MIMES)


class UnsupportedDocumentError(ValueError):
    pass


class TooManyPagesError(ValueError):
    def __init__(self, pages: int, limit: int) -> None:
        super().__init__(f"the document has {pages} pages; the limit is {limit}")
        self.pages = pages
        self.limit = limit


class DocumentParseError(ValueError):
    pass


def check_zip_package(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise DocumentParseError("the DOCX file is not a valid zip package") from exc
    if len(entries) > MAX_DOCX_ENTRIES:
        raise DocumentParseError("the DOCX package has too many entries")
    uncompressed = sum(entry.file_size for entry in entries)
    compressed = sum(entry.compress_size for entry in entries)
    if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise DocumentParseError("the DOCX package expands beyond the allowed size")
    if compressed > 0 and uncompressed / compressed > MAX_DOCX_COMPRESSION_RATIO:
        raise DocumentParseError("the DOCX package has a suspicious compression ratio")


def detect_format(data: bytes, mime: str | None, filename: str | None = None) -> str:
    mime = (mime or "").split(";")[0].strip().lower()
    if mime in PDF_MIMES:
        return "pdf"
    if mime in DOCX_MIMES:
        return "docx"
    if mime in MARKDOWN_MIMES:
        return "md"
    if mime in HTML_MIMES:
        return "html"
    if mime in IMAGE_MIMES:
        return "image"
    if filename:
        ext = PurePath(filename).suffix.lower()
        if ext in EXTENSIONS:
            return EXTENSIONS[ext]
    head = data[:16]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04") and b"word/" in data[:4096] + data[-65536:]:
        return "docx"
    if head.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"II*\x00", b"MM\x00*", b"BM")) or head[8:12] == b"WEBP":
        return "image"
    if mime in TEXT_MIMES or mime.startswith("text/") or not mime or mime == "application/octet-stream":
        sniff = data[:2048].lstrip().lower()
        if sniff.startswith((b"<!doctype html", b"<html")):
            return "html"
        return "txt"
    raise UnsupportedDocumentError(f"unsupported document type: {mime or 'unknown'}")


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    cp1251 = data.decode("cp1251", errors="replace")
    cyrillic = sum(1 for ch in cp1251 if "Ѐ" <= ch <= "ӿ")
    letters = sum(1 for ch in cp1251 if ch.isalpha()) or 1
    if cyrillic / letters > 0.3:
        return cp1251
    return data.decode("latin-1")


class DocumentParser:
    def __init__(self, ocr: OcrEngine) -> None:
        self._ocr = ocr
        self._handlers: dict[str, Callable[[bytes], ParsedDocument]] = {
            "pdf": self.parse_pdf,
            "docx": self.parse_docx,
            "txt": self.parse_text,
            "md": self.parse_markdown,
            "html": self.parse_html,
            "image": self.parse_image,
        }

    def parse(
        self, data: bytes, mime: str | None, filename: str | None = None, max_pages: int | None = None
    ) -> ParsedDocument:
        if not data:
            raise DocumentParseError("the file is empty")
        fmt = detect_format(data, mime, filename)
        if max_pages is not None:
            self._check_page_limit(fmt, data, max_pages)
        try:
            parsed = self._handlers[fmt](data)
        except (UnsupportedDocumentError, DocumentParseError, OcrUnavailableError, TooManyPagesError):
            raise
        except Exception as exc:
            raise DocumentParseError(f"cannot parse {fmt} document: {exc}") from exc
        parsed.blocks = [b for b in parsed.blocks if b.text.strip()]
        return parsed

    @staticmethod
    def _check_page_limit(fmt: str, data: bytes, max_pages: int) -> None:
        count = 0
        try:
            if fmt == "pdf":
                from pypdf import PdfReader

                count = len(PdfReader(io.BytesIO(data)).pages)
            elif fmt == "image":
                from PIL import Image

                with Image.open(io.BytesIO(data)) as image:
                    count = int(getattr(image, "n_frames", 1))
        except Exception:
            return
        if count > max_pages:
            raise TooManyPagesError(count, max_pages)

    def parse_pdf(self, data: bytes) -> ParsedDocument:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentParseError("the PDF is password-protected")
        page_count = len(reader.pages)
        doc = ParsedDocument(blocks=[], pages=page_count)
        ocr_pages: list[int] = []
        texts: list[str] = []
        for index, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                log.warning("pdf_page_text_failed", page=index + 1, error=str(exc))
                text = ""
            texts.append(text)
            if len(text.strip()) < MIN_PAGE_TEXT_CHARS:
                ocr_pages.append(index)

        if ocr_pages:
            self._ocr_pdf_pages(data, ocr_pages, texts, doc)

        for index, text in enumerate(texts):
            doc.blocks.extend(lines_to_blocks(text, page=index + 1))
        if not doc.blocks:
            raise DocumentParseError("no text could be extracted from the PDF")
        return doc

    def _ocr_pdf_pages(self, data: bytes, pages: list[int], texts: list[str], doc: ParsedDocument) -> None:
        if not self._ocr.available():
            doc.warnings.append(f"{len(pages)} page(s) have no text layer and OCR is unavailable")
            return
        self._note_missing_languages(doc)
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(data)
        try:
            for index in pages[:MAX_OCR_PAGES]:
                page = pdf[index]
                try:
                    image = page.render(scale=self._ocr.dpi / 72).to_pil()
                    ocr_text = self._ocr.image_to_text(image)
                finally:
                    page.close()
                if len(ocr_text.strip()) > len(texts[index].strip()):
                    texts[index] = ocr_text
                    doc.ocr_used = True
        finally:
            pdf.close()
        if len(pages) > MAX_OCR_PAGES:
            doc.warnings.append(f"OCR limited to the first {MAX_OCR_PAGES} pages without a text layer")

    def _note_missing_languages(self, doc: ParsedDocument) -> None:
        missing = self._ocr.missing_languages
        if missing:
            doc.warnings.append(f"OCR language data missing: {', '.join(missing)}; used {self._ocr.languages()}")

    def parse_image(self, data: bytes) -> ParsedDocument:
        from PIL import Image, ImageSequence

        if not self._ocr.available():
            raise OcrUnavailableError("images need OCR, but Tesseract is not available")
        doc = ParsedDocument(blocks=[], pages=0, ocr_used=True)
        self._note_missing_languages(doc)
        Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
        with Image.open(io.BytesIO(data)) as image:
            for index, frame in enumerate(ImageSequence.Iterator(image)):
                if index >= MAX_OCR_PAGES:
                    break
                text = self._ocr.image_to_text(frame.convert("RGB"))
                doc.pages = index + 1
                doc.blocks.extend(lines_to_blocks(text, page=index + 1))
        if not doc.blocks:
            raise DocumentParseError("OCR found no text in the image")
        return doc

    def parse_docx(self, data: bytes) -> ParsedDocument:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        check_zip_package(data)
        document = docx.Document(io.BytesIO(data))
        blocks: list[Block] = []
        body = document.element.body
        for child in body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                block = _docx_paragraph_block(Paragraph(child, document))
                if block:
                    blocks.append(block)
            elif tag == "tbl":
                rows = []
                for row in Table(child, document).rows:
                    cells: list[str] = []
                    for cell in row.cells:
                        value = " ".join(cell.text.split())
                        if value and (not cells or cells[-1] != value):
                            cells.append(value)
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    blocks.append(Block("\n".join(rows), "table"))
        return ParsedDocument(blocks=blocks, pages=None)

    def parse_text(self, data: bytes) -> ParsedDocument:
        return ParsedDocument(blocks=lines_to_blocks(decode_text(data)), pages=None)

    def parse_markdown(self, data: bytes) -> ParsedDocument:
        from markdown_it import MarkdownIt

        md = MarkdownIt("commonmark").enable("table")
        tokens = md.parse(decode_text(data))
        blocks: list[Block] = []
        list_depth = 0
        row: list[str] = []
        table_rows: list[str] = []
        pending_heading: int | None = None
        for i, token in enumerate(tokens):
            ttype = token.type
            if ttype == "heading_open":
                pending_heading = int(token.tag[1])
            elif ttype in ("bullet_list_open", "ordered_list_open"):
                list_depth += 1
            elif ttype in ("bullet_list_close", "ordered_list_close"):
                list_depth -= 1
            elif ttype == "tr_open":
                row = []
            elif ttype == "tr_close":
                if row:
                    table_rows.append(" | ".join(row))
            elif ttype == "table_close":
                if table_rows:
                    blocks.append(Block("\n".join(table_rows), "table"))
                table_rows = []
            elif ttype in ("fence", "code_block"):
                if token.content.strip():
                    blocks.append(Block(token.content.rstrip(), "code"))
            elif ttype == "inline":
                text = _inline_text(token)
                prev = tokens[i - 1].type if i else ""
                if prev in ("th_open", "td_open"):
                    row.append(text)
                elif pending_heading is not None:
                    blocks.append(Block(text, "heading", level=pending_heading))
                    pending_heading = None
                elif text:
                    blocks.append(Block(text, "list" if list_depth else "paragraph"))
        return ParsedDocument(blocks=blocks, pages=None)

    def parse_html(self, data: bytes) -> ParsedDocument:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(decode_text(data), "html.parser")
        for tag in soup(["script", "style", "noscript", "template", "nav", "footer", "aside", "form", "svg", "iframe"]):
            tag.decompose()
        root = soup.body or soup
        blocks: list[Block] = []
        for element in root.find_all(True):
            name = element.name
            if name != "table" and element.find_parent("table") is not None:
                continue
            if name == "table":
                rows = []
                for tr in element.find_all("tr"):
                    cells = [" ".join(c.get_text(" ").split()) for c in tr.find_all(["th", "td"])]
                    cells = [c for c in cells if c]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    blocks.append(Block("\n".join(rows), "table"))
            elif re.fullmatch(r"h[1-6]", name):
                text = " ".join(element.get_text(" ").split())
                if text:
                    blocks.append(Block(text, "heading", level=int(name[1])))
            elif name in _HTML_BLOCKS and not element.find(_HTML_BLOCKS_AND_HEADINGS):
                text = " ".join(element.get_text(" ").split())
                if text:
                    kind: BlockKind = (
                        "list"
                        if name == "li" or element.find_parent("li")
                        else ("code" if name == "pre" else "paragraph")
                    )
                    blocks.append(Block(text, kind))
        return ParsedDocument(blocks=blocks, pages=None)


_HTML_BLOCKS = [
    "p",
    "li",
    "pre",
    "blockquote",
    "dt",
    "dd",
    "figcaption",
    "caption",
    "div",
    "section",
    "article",
    "main",
    "td",
]
_HTML_BLOCKS_AND_HEADINGS = [*_HTML_BLOCKS, "h1", "h2", "h3", "h4", "h5", "h6", "table", "ul", "ol"]
_DOCX_HEADING_STYLE = re.compile(r"(?:heading|заголовок)\s*(\d)", re.IGNORECASE)


def _docx_paragraph_block(paragraph: Any) -> Block | None:
    text = " ".join(paragraph.text.split())
    if not text:
        return None
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    match = _DOCX_HEADING_STYLE.search(style_name)
    if match:
        return Block(text, "heading", level=max(1, min(6, int(match.group(1)))))
    if style_name.lower() in ("title", "название", "заголовок"):
        return Block(text, "heading", level=1)
    p_pr = paragraph._p.pPr
    if (p_pr is not None and p_pr.numPr is not None) or "list" in style_name.lower():
        return Block(text, "list")
    level = heading_level(text) if len(text) < 80 and _is_bold(paragraph) else 0
    return Block(text, "heading", level=level) if level else Block(text, "paragraph")


def _is_bold(paragraph: Any) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    return bool(runs) and all(bool(r.bold) for r in runs)


def _inline_text(token: Any) -> str:
    parts: list[str] = []
    for child in token.children or []:
        if child.type in ("text", "code_inline"):
            parts.append(child.content)
        elif child.type in ("softbreak", "hardbreak"):
            parts.append(" ")
        elif child.type == "image":
            parts.append(child.content)
    return " ".join("".join(parts).split())
